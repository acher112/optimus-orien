"""Shared export service for Optimus Orien.
Generates CSV (with UTF-8 BOM), Excel (via openpyxl), and PDF (via reportlab)
with formula-injection protection, 50,000 row limits, audit logging, and security headers.
"""
import io
import re
import csv
import codecs
from datetime import date, datetime
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.response import Response

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import pagesizes, colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

from core.services.audit import record_event

MAX_EXPORT_ROWS = 50000
INJECTION_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def sanitize_cell_value(val):
    """Spreadsheet-injection neutralization:
    Any text cell starting with =, +, -, @, tab or CR is prefixed with ' in CSV and XLSX.
    """
    if val is None:
        return ''
    if isinstance(val, (int, float, Decimal, bool)):
        return val
    s = str(val)
    if s and s[0] in INJECTION_PREFIXES:
        return "'" + s
    return s


def sanitize_filename(name):
    """Sanitize filename to prevent directory traversal or invalid characters."""
    s = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', str(name))
    return s.strip('._') or 'export'


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas for exact 'Page X of Y' numbering and running header/footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor('#64748B'))
        page_w, page_h = pagesizes.landscape(pagesizes.A4)

        # Footer
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(page_w - 28, 16, footer_text)
        self.drawString(28, 16, "Optimus Orien Educational Platform • Confidential & Authorised Use Only")

        # Running top rule
        self.setStrokeColor(colors.HexColor('#E2E8F0'))
        self.setLineWidth(0.5)
        self.line(28, page_h - 24, page_w - 28, page_h - 24)
        self.restoreState()


def export_response(request, filename, title, columns, rows, fmt, meta=None):
    """Generate and return an export download response.
    - fmt: 'csv', 'xlsx', 'pdf'
    - columns: list of dicts [{'key': 'admission_number', 'label': 'Admission #', 'type': 'text'}]
    - rows: list of dicts matching column keys
    - meta: optional dict of context (school_name, date_range, filters, etc.)
    """
    if len(rows) > MAX_EXPORT_ROWS:
        return Response({
            'detail': f'Export exceeds maximum allowed limit of {MAX_EXPORT_ROWS:,} rows. Please refine your search filters.'
        }, status=413)

    fmt = str(fmt or 'csv').lower().strip()
    if fmt not in ('csv', 'xlsx', 'pdf'):
        fmt = 'csv'

    meta = meta or {}
    user = getattr(request, 'user', None)
    today_str = date.today().strftime('%Y-%m-%d')
    safe_name = sanitize_filename(filename)
    download_filename = f"{safe_name}-{today_str}.{fmt}"

    # Extract headers and keys
    col_defs = []
    for c in columns:
        if isinstance(c, dict):
            col_defs.append({
                'key': c.get('key') or c.get('name'),
                'label': c.get('label') or c.get('title') or str(c.get('key')),
                'type': c.get('type', 'text')
            })
        elif isinstance(c, (list, tuple)):
            col_defs.append({'label': str(c[0]), 'key': str(c[1]), 'type': 'text'})
        else:
            col_defs.append({'label': str(c), 'key': str(c), 'type': 'text'})

    # Audit export event
    try:
        record_event(
            actor=user,
            action='export',
            obj=None,
            extras={
                'module': safe_name,
                'format': fmt,
                'rows': len(rows),
                'filters': str(meta.get('filters', ''))
            }
        )
    except Exception:
        pass

    # ------------------------------------------------ CSV GENERATION
    if fmt == 'csv':
        buffer = io.BytesIO()
        # Write UTF-8 BOM so Microsoft Excel correctly reads non-Latin text (Urdu, Arabic, etc.)
        buffer.write(codecs.BOM_UTF8)
        text_stream = io.StringIO()
        writer = csv.writer(text_stream, quoting=csv.QUOTE_MINIMAL)

        # Header row
        writer.writerow([c['label'] for c in col_defs])

        # Data rows
        for r in rows:
            row_vals = []
            for c in col_defs:
                v = r.get(c['key']) if isinstance(r, dict) else getattr(r, c['key'], '')
                if isinstance(v, Decimal):
                    row_vals.append(f"{v:.2f}")
                elif v is None:
                    row_vals.append('')
                else:
                    row_vals.append(sanitize_cell_value(v))
            writer.writerow(row_vals)

        buffer.write(text_stream.getvalue().encode('utf-8'))
        response = HttpResponse(buffer.getvalue(), content_type='text/csv; charset=utf-8')

    # ------------------------------------------------ XLSX GENERATION
    elif fmt == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (title or 'Export')[:31]

        # Freeze header row
        ws.freeze_panes = 'A2'

        # Styles
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        header_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        data_font = Font(name='Segoe UI', size=10)
        data_border = Border(
            bottom=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0')
        )
        currency_format = '#,##0.00'
        date_format = 'yyyy-mm-dd'

        # Write header
        ws.row_dimensions[1].height = 26
        for col_idx, col in enumerate(col_defs, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col['label'])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Write data rows
        for row_idx, r in enumerate(rows, start=2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx, col in enumerate(col_defs, start=1):
                val = r.get(col['key']) if isinstance(r, dict) else getattr(r, col['key'], '')
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = data_border

                if isinstance(val, Decimal):
                    cell.value = float(val)
                    cell.number_format = currency_format
                    cell.alignment = Alignment(horizontal='right')
                elif isinstance(val, (int, float)):
                    cell.value = val
                    cell.alignment = Alignment(horizontal='right')
                elif isinstance(val, (date, datetime)):
                    cell.value = val
                    cell.number_format = date_format
                    cell.alignment = Alignment(horizontal='center')
                else:
                    cell.value = sanitize_cell_value(val)
                    cell.alignment = Alignment(horizontal='left')

        # Auto-fit column widths
        for col_idx, col in enumerate(col_defs, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col['label']))
            for row_idx in range(2, min(len(rows) + 2, 200)):
                cv = ws.cell(row=row_idx, column=col_idx).value
                if cv is not None:
                    max_len = max(max_len, len(str(cv)))
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    # ------------------------------------------------ PDF GENERATION
    else:
        buffer = io.BytesIO()
        page_w, page_h = pagesizes.landscape(pagesizes.A4)
        doc = SimpleDocTemplate(
            buffer,
            pagesize=(page_w, page_h),
            leftMargin=28,
            rightMargin=28,
            topMargin=32,
            bottomMargin=32
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=4
        )
        meta_style = ParagraphStyle(
            'DocMeta',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=12
        )
        th_style = ParagraphStyle(
            'TH',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11,
            textColor=colors.white,
            alignment=0
        )
        td_style = ParagraphStyle(
            'TD',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#1E293B')
        )

        story = []

        # Document Header
        story.append(Paragraph(title or safe_name.replace('_', ' ').title(), title_style))
        campus_info = meta.get('school_name') or meta.get('campus', '')
        gen_by = user.get_full_name() or user.username if user else 'System'
        meta_line = f"Campus: <b>{campus_info or 'All Campuses'}</b> • Generated by: <b>{gen_by}</b> • Date: <b>{today_str}</b>"
        if meta.get('filters'):
            meta_line += f" • Filters: <i>{meta['filters']}</i>"
        story.append(Paragraph(meta_line, meta_style))
        story.append(Spacer(1, 6))

        # Build table matrix
        table_data = []
        table_data.append([Paragraph(c['label'], th_style) for c in col_defs])

        for r in rows:
            row_cells = []
            for c in col_defs:
                v = r.get(c['key']) if isinstance(r, dict) else getattr(r, c['key'], '')
                if isinstance(v, Decimal):
                    display_text = f"Rs {v:,.2f}"
                elif v is None or v == '':
                    display_text = '—'
                else:
                    display_text = str(v)
                # Escape XML/HTML tags in PDF text
                display_text = display_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                row_cells.append(Paragraph(display_text, td_style))
            table_data.append(row_cells)

        # Distribute available page width across columns
        available_w = page_w - 56
        col_count = max(len(col_defs), 1)
        col_widths = [available_w / col_count] * col_count

        pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        story.append(pdf_table)

        doc.build(story, canvasmaker=NumberedCanvas)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')

    # Security & Download Headers
    response['Content-Disposition'] = f'attachment; filename="{download_filename}"'
    response['X-Content-Type-Options'] = 'nosniff'
    response['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    response['Pragma'] = 'no-cache'
    return response


def handle_export(request, data_or_qs, filename, title, columns, meta=None):
    """Convenience helper for API endpoints.
    If ?format=csv|xlsx|pdf is in query params, builds row list and returns export response.
    Otherwise returns None so the calling view proceeds with regular JSON response.
    """
    fmt = request.query_params.get('format') if hasattr(request, 'query_params') else request.GET.get('format')
    if not fmt or str(fmt).lower() not in ('csv', 'xlsx', 'pdf'):
        return None

    # Resolve queryset or iterable
    rows = []
    if hasattr(data_or_qs, 'all') and callable(getattr(data_or_qs, 'all')):
        # Model queryset
        for obj in data_or_qs:
            row_dict = {}
            for col in columns:
                k = col.get('key') if isinstance(col, dict) else (col[1] if isinstance(col, tuple) else str(col))
                val = getattr(obj, k, None)
                if callable(val):
                    val = val()
                row_dict[k] = val
            rows.append(row_dict)
    elif isinstance(data_or_qs, list):
        rows = data_or_qs
    else:
        rows = list(data_or_qs)

    return export_response(request, filename, title, columns, rows, fmt, meta=meta)
